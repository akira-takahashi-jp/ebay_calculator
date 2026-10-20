"""利益計算表（xlsx）から料金表・カテゴリー手数料を抜き出して data.js を生成する。

使い方:
    python3 tools/extract_data.py 利益計算表配布用2.xlsx data.js
（openpyxl が必要）
"""
import json
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as ci
from openpyxl.utils import get_column_letter


def num(v):
    assert isinstance(v, (int, float)), v
    return int(v) if float(v).is_integer() else v


def pairs(ws, wcol, pcol, r1, r2, mult=1):
    return [[num(ws[f"{wcol}{r}"].value * mult), num(ws[f"{pcol}{r}"].value)] for r in range(r1, r2 + 1)]


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    d = {}

    s = wb["送料"]
    d["airUS"] = pairs(s, "I", "J", 2, 21)
    d["airEU"] = pairs(s, "B", "C", 2, 21)
    d["ems"] = pairs(s, "B", "C", 23, 64)
    d["print"] = pairs(s, "B", "C", 66, 89)
    d["letter"] = pairs(s, "L", "M", 2, 7)

    e = wb["エコノミー送料"]
    d["ecoUS"] = pairs(e, "B", "D", 10, 75, 1000)
    d["ecoUSOther"] = pairs(e, "B", "D", 82, 127, 1000)
    cols = range(ci("L"), ci("AL") + 1)
    d["ecoEU"] = {
        "codes": [e.cell(9, c).value for c in cols],
        "rows": [[num(e[f"J{r}"].value * 1000)] + [num(e.cell(r, c).value) for c in cols] for r in range(10, 66)],
    }

    f = wb["FedEx 送料"]
    d["fedex"] = {
        "zones": [f.cell(18, c).value for c in range(4, 9)],
        "rows": [[num(f[f"B{r}"].value * 1000)] + [num(f.cell(r, c).value) for c in range(4, 9)] for r in range(21, 109)],
    }

    h = wb["DHL　送料"]
    d["dhl"] = [[num(h[f"B{r}"].value * 1000), num(h[f"D{r}"].value), num(h[f"E{r}"].value)] for r in range(21, 195)]

    st = wb["設定"]
    d["euCountries"] = [
        {"name": st[f"F{r}"].value, "code": st[f"G{r}"].value, "ficp": st[f"H{r}"].value, "dhl": num(st[f"I{r}"].value)}
        for r in range(50, 77)
    ]

    c = wb["カテゴリー手数料一覧"]
    cats = []
    for col in range(ci("B"), ci("U") + 1):
        g = lambda r: c.cell(r, col).value  # noqa: E731
        ids = []
        for r in range(8, 2103):
            v = g(r)
            if v is None:
                continue
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            ids.append(str(v).strip())
        cats.append({
            "col": get_column_letter(col),
            "base": g(3),
            "t1": g(4),
            "r1": g(5),
            "t2": g(6),
            "r2": g(7),
            "ids": ids,
        })
    d["categories"] = cats

    with open(dst, "w", encoding="utf-8") as out:
        out.write("// tools/extract_data.py で xlsx から自動生成。手で編集しないこと。\n")
        out.write("window.RATE_DATA = ")
        json.dump(d, out, ensure_ascii=False, separators=(",", ":"))
        out.write(";\n")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
